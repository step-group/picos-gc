"""Repeat campaign (REPETICIONES_22SEP2026) -> mass fractions, old vs new tie-lines.

The batch ran on a new 17-min method and mixes every block in one folder, so it gets
its own fixed retention-time map keyed on the sample-name prefix (C4_T1 -> block C,
BIN3_B2 -> block C's water-solvent edge, 2PE_T1 -> the water-2PE binary). Its areas,
with the masses and organic KF typed into out/repeat_entry.xlsx, replace the
campaign-1 vials of the same tubes in an in-memory copy of the filled workbook; a
re-equilibrated tube is a new tie-line, so all its vials are replaced, both phases.
fill_ternarios' own chain then gives the tie-lines (slopes: CC_MF, the May
calibration) and audit_repeats re-judges them. The master workbooks are not written.

Run from the repo root: uv run python repeat_results.py
  -> out/REPETICIONES_22SEP2026/{areas.csv, <sample>_labeled.png}
     out/repeat_tielines.csv   old (campaign 1) and new points of every repeated tube
     out/repeat_list_after.csv audit_repeats verdicts with the repeats in
     out/repeat_compare/<block>.png
"""

from __future__ import annotations

import csv
import math
import re
import shutil
from pathlib import Path

import matplotlib.pyplot as plt
import openpyxl

from audit_repeats import COLS, audit, ternary_recs
from fill_ternarios import (
    _BIN_ROWS,
    BIN_TO_BLOCK,
    ORGANIC_WATER_MAX,
    WB_OUT,
    _vial_rows,
    binary_tieline_rows,
    canon,
    cc_mf_slopes,
    results_rows,
)
from label_terpenos import (
    IPA_WINDOW,
    PARAMS,
    in_window,
    load_chrom,
    plot_labeled,
    safe,
    sample_label,
)
from picos_gc.detector import detect_peaks
from picos_gc.integrator import integrate_all_peaks

_ROOT = Path(__file__).resolve().parent
BATCH = _ROOT / "FLECK_TERPENOS2026" / "REPETICIONES_22SEP2026"
ENTRY = _ROOT / "out" / "repeat_entry.xlsx"
CC_MF = _ROOT / "CC_MF.xlsx"
OUT = _ROOT / "out"

# ponytail: retention times of THIS batch's method (read off its chromatograms); a
# batch on another method needs its own map.
REPEAT_TR = {
    "DL-Camphor": 8.45,
    "L-Carvone": 9.67,
    "2PE": 10.72,
    "Thymol": 13.39,
    "Eugenol": 13.79,  # drifts 13.69-13.89 with load; never shares a block with thymol's
    "Carvacrol": 14.03,  # neighbour carvacrol, and D's thymol sits 0.3 min below it
}
RT_TOL = 0.25
BLOCK_TERPS = {  # label_terpenos.BATCH_TERPENES, keyed on the block letter
    "A": ("L-Carvone", "Thymol"),
    "C": ("Thymol", "Carvacrol"),
    "D": ("Thymol", "Eugenol"),
    "E": ("DL-Camphor", "Thymol"),
    "F": ("DL-Camphor", "L-Carvone"),
    "H": ("DL-Camphor", "Carvacrol"),
    "I": ("DL-Camphor", "Eugenol"),
}
CODE = re.compile(r"^(.+?)[-_ ]?([TB])([12])$")


def parse_code(code: str) -> tuple[str, str, int]:
    """'C4_T1' / 'C4-T1' -> ('C4', 'T', 1); 'BIN3_B2' -> ('BIN3', 'B', 2)."""
    m = CODE.match(code.strip())
    if not m:
        raise ValueError(f"not a vial code: {code!r}")
    return m.group(1), m.group(2), int(m.group(3))


def block_of(tube: str) -> str:
    """Block letter of a tube; 'BIN<n>' is block BIN_TO_BLOCK[n]'s edge, '2PE' its own."""
    if tube.startswith("BIN"):
        return BIN_TO_BLOCK[int(tube[3:])]
    return "2PE" if tube == "2PE" else tube[0]


def compounds(tube: str) -> tuple[str, ...]:
    """The compounds a vial of this tube can hold (binary edges carry no 2PE)."""
    if tube == "2PE":
        return ("2PE",)
    terps = BLOCK_TERPS[block_of(tube)]
    return terps if tube.startswith("BIN") else ("2PE", *terps)


def classify_repeat(peaks, names) -> list[tuple[str, object]]:
    """label_terpenos.classify with fixed centres: each peak goes to the nearest of
    *names* within RT_TOL; each compound keeps its largest peak, the rest are 'other'."""
    claims, out = {}, []
    for p in peaks:
        if in_window(p.time_min, IPA_WINDOW):
            claims.setdefault("IPA", []).append(p)
            continue
        d, nm = min((abs(p.time_min - REPEAT_TR[n]), n) for n in names)
        if d <= RT_TOL:
            claims.setdefault(nm, []).append(p)
        else:
            out.append(("other", p))
    for nm, plist in claims.items():
        plist.sort(key=lambda p: p.area_mV_min, reverse=True)
        out.append((nm, plist[0]))
        out.extend(("other", p) for p in plist[1:])
    out.sort(key=lambda ap: ap[1].time_min)
    return out


def integrate() -> dict[str, dict[str, float]]:
    """{vial code: {canon compound: area}} for the batch; writes areas.csv + PNGs."""
    out_dir = OUT / BATCH.name
    if out_dir.exists():
        shutil.rmtree(out_dir)
    out_dir.mkdir(parents=True)
    areas, rows = {}, []
    for fp in sorted(BATCH.glob("*.gcd")):
        ch = load_chrom(fp)
        name = sample_label(ch)
        if name == "BLANK":
            continue
        tube, ph, rep = parse_code(name)
        code = f"{tube}-{ph}{rep}"  # the entry sheet's spelling
        names = compounds(tube)
        assigns = classify_repeat(integrate_all_peaks(ch, detect_peaks(ch, PARAMS)), names)
        found = {nm: p for nm, p in assigns if nm in names}
        areas[code] = {canon(nm): p.area_mV_min for nm, p in found.items()}
        missing = [n for n in names if n not in found]
        rows.append(
            [code, fp.name]
            + [f"{found[n].time_min:.3f}" if n in found else "" for n in REPEAT_TR]
            + [f"{found[n].area_mV_min:.4f}" if n in found else "" for n in REPEAT_TR]
            + [";".join(missing)]
        )
        plot_labeled(ch, assigns, f"{code}  [{fp.name}]", out_dir / f"{safe(code)}_labeled.png")
    with (out_dir / "areas.csv").open("w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(
            ["vial", "file"]
            + [f"{canon(n)}_tR" for n in REPEAT_TR]
            + [f"{canon(n)}_area" for n in REPEAT_TR]
            + ["missing"]
        )
        w.writerows(rows)
    return areas


def read_entry(path: Path = ENTRY) -> dict[str, dict[str, float | None]]:
    """{vial code: master-column values} from the typed entry sheet (row 9 on):
    E/G/I masses -> D/F/H, M/N KF -> U/V. A vial never weighed (2PE-B1/B2: KF only) gets
    None masses."""
    ws = openpyxl.load_workbook(path, data_only=True)["Entry"]
    out = {}
    for r in ws.iter_rows(min_row=9, values_only=True):
        code = r[1]
        if not code:
            continue
        num = [v if isinstance(v, int | float) and v else None for v in r]
        out[code] = {"D": num[4], "F": num[6], "H": num[8], "U": num[12], "V": num[13]}
    return out


def _row_of(ws, tube: str, ph: str, rep: int) -> int:
    if tube.startswith("BIN"):
        return _BIN_ROWS[(ph, rep)]
    sysnum = int(tube[1:])
    return next(r for r, s, p, k, _ in _vial_rows(ws) if (s, p, k) == (sysnum, ph, rep))


def patch(wb, areas: dict, entry: dict) -> list[str]:
    """Overwrite each repeat vial's row (masses, KF, areas) in the data_only workbook.
    Returns the tubes patched. The 2PE-water binary has no sheet; see two_pe_endpoint."""
    tubes = []
    for code, cells in entry.items():
        tube, ph, rep = parse_code(code)
        if tube == "2PE" or None in (cells["D"], cells["F"], cells["H"]):
            continue
        if code not in areas:
            print(f"  WARNING {code}: weighed but no chromatogram; campaign-1 vial kept")
            continue
        ws = wb[f"Bloque {block_of(tube)}"]
        r = _row_of(ws, tube, ph, rep)
        a = areas[code]
        for col, v in cells.items():
            ws[f"{col}{r}"] = v
        ws[f"L{r}"] = None if tube.startswith("BIN") else a.get("2pe")
        ws[f"M{r}"] = a.get(canon(ws["M3"].value))
        ws[f"N{r}"] = a.get(canon(ws["N3"].value))
        if tube not in tubes:
            tubes.append(tube)
    return tubes


def two_pe_endpoint(areas: dict, entry: dict, f2: float) -> list[dict]:
    """The water-2PE binary: organic 2PE = 1 - KF (nothing else in it), aqueous 2PE from
    the GC, water by difference. Same arithmetic as fill_ternarios.vial_fractions."""
    org = [c / 100 for k, e in entry.items() if k.startswith("2PE-") for c in (e["U"], e["V"])
           if c is not None]  # fmt: skip
    aq = [
        areas[k]["2pe"] / f2 * (e["H"] - e["D"]) / (e["F"] - e["D"]) / 100
        for k, e in entry.items()
        if k.startswith("2PE-") and "2pe" in areas.get(k, {}) and None not in (e["D"], e["F"])
    ]
    out = []
    if org:
        w = sum(org) / len(org)
        out.append(_pt("2PE", "2PE-bin", "organic", 1 - w, 0.0, 0.0, w, "kf"))
    if aq:
        s = sum(aq) / len(aq)
        out.append(_pt("2PE", "2PE-bin", "aqueous", s, 0.0, 0.0, 1 - s, "bydiff"))
    return out


def _pt(block, system, phase, s, a, b, w, src, hba="", hbd="", closure="", flags=""):
    return {"block": block, "system": system, "phase": phase, "w_2pe": s, "hba": hba,
            "w_hba": a, "hbd": hbd, "w_hbd": b, "water": w, "closure": closure,
            "flags": flags, "water_src": src}  # fmt: skip


def tielines(wb) -> list[dict]:
    """Every point of a data_only workbook through fill_ternarios' chain, as dicts."""
    out = []
    for sheet in wb.sheetnames:
        ws, block = wb[sheet], sheet.split()[-1]
        for r in results_rows(ws, block, ternary_recs(ws)):
            if r[3]:
                out.append(_pt(r[0], r[1], r[2], *(float(r[i]) for i in (3, 5, 7, 8)), r[12],
                               r[4], r[6], r[9], r[11]))  # fmt: skip
    for b, phase, hba, x, hbd, y, z, src in binary_tieline_rows(wb):
        out.append(_pt(b, f"{b}-bin", phase, 0.0, float(x), float(y), float(z), src, hba, hbd))
    return out


def _system(tube: str) -> str:
    return f"{block_of(tube)}-bin" if tube.startswith(("BIN", "2PE")) else tube


# --- plot: pseudo-ternary 2PE / DES (HBA+HBD) / water, mass fractions -------------
def _xy(p):
    des = p["w_hba"] + p["w_hbd"]
    return des + 0.5 * p["w_2pe"], math.sqrt(3) / 2 * p["w_2pe"]


def _pairs(points):
    by = {}
    for p in points:
        by.setdefault(p["system"], []).append(p)
    return {s: ps for s, ps in by.items() if len(ps) == 2}


def plot_block(block, old, new, failed, two_pe, path):
    fig, (ax, az) = plt.subplots(1, 2, figsize=(11, 5))
    ax.plot([0, 1, 0.5, 0], [0, 0, math.sqrt(3) / 2, 0], color="black", lw=0.8)
    ax.text(0, -0.04, "water", ha="center", va="top")
    ax.text(1, -0.04, "DES", ha="center", va="top")
    ax.text(0.5, math.sqrt(3) / 2 + 0.02, "2PE", ha="center")
    labelled = set()

    def draw(ps, style, label):
        xs, ys = zip(*(_xy(p) for p in ps), strict=True)
        ax.plot(xs, ys, **style, label=None if label in labelled else label)
        labelled.add(label)
        for p in ps:
            if p["water"] > ORGANIC_WATER_MAX:
                az.loglog(max(p["w_hba"] + p["w_hbd"], 1e-5), max(p["w_2pe"], 1e-5),
                          style.get("marker") or "o", color=style["color"], ms=6,
                          mfc="none" if style["ls"] == "--" else style["color"])  # fmt: skip
                az.annotate(p["system"], (max(p["w_hba"] + p["w_hbd"], 1e-5),
                            max(p["w_2pe"], 1e-5)), fontsize=7, color=style["color"],
                            xytext=(3, 3), textcoords="offset points")  # fmt: skip

    for s, ps in _pairs(old).items():
        if s in failed:
            draw(ps, {"color": "tab:red", "ls": "--", "marker": "x", "lw": 1}, "campaign 1, failed")
        elif s not in {p["system"] for p in new}:
            draw(ps, {"color": "0.6", "ls": "-", "marker": "o", "ms": 3, "lw": 0.8}, "campaign 1")
        else:
            draw(ps, {"color": "tab:orange", "ls": "--", "marker": "x", "lw": 1},
                 "campaign 1, re-run")  # fmt: skip
    for ps in _pairs(new).values():
        draw(ps, {"color": "tab:blue", "ls": "-", "marker": "o", "ms": 4, "lw": 1.4}, "repeat")
    if len(two_pe) == 2:
        draw(two_pe, {"color": "tab:green", "ls": "-", "marker": "s", "ms": 4, "lw": 1.4},
             "water-2PE binary, repeat")  # fmt: skip
    ax.set_aspect("equal")
    ax.axis("off")
    ax.legend(fontsize=8, loc="upper right")
    ax.set_title(f"Block {block} (mass fractions)")
    az.set_xlabel("w(HBA + HBD), aqueous phase (0 drawn at 1e-5)")
    az.set_ylabel("w(2PE), aqueous phase (0 drawn at 1e-5)")
    az.set_title("aqueous endpoints")
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


FIELDS = ["campaign", "block", "system", "phase", "w_2pe", "hba", "w_hba", "hbd", "w_hbd",
          "water", "closure", "flags", "water_src"]  # fmt: skip


def main() -> None:
    areas = integrate()
    entry = read_entry()
    wb_old = openpyxl.load_workbook(WB_OUT, data_only=True)
    wb_new = openpyxl.load_workbook(WB_OUT, data_only=True)
    tubes = patch(wb_new, areas, entry)
    systems = {_system(t) for t in tubes}
    old, new = tielines(wb_old), tielines(wb_new)
    two_pe = two_pe_endpoint(areas, entry, cc_mf_slopes(CC_MF)["2pe"])

    rows = [{"campaign": "1", **p} for p in old if p["system"] in systems]
    rows += [{"campaign": "repeat", **p} for p in new if p["system"] in systems]
    rows += [{"campaign": "repeat", **p} for p in two_pe]
    rows.sort(key=lambda r: (r["system"], r["phase"], r["campaign"]))
    with (OUT / "repeat_tielines.csv").open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS)
        w.writeheader()
        for r in rows:
            w.writerow({k: f"{v:.5f}" if isinstance(v, float) else v for k, v in r.items()})

    before = {(r["system"], r["phase"]): r for r in audit(wb_old)}
    after = audit(wb_new)
    with (OUT / "repeat_list_after.csv").open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=COLS)
        w.writeheader()
        w.writerows(after)
    print(f"\n{'point':<9}{'phase':<10}{'before':<50}after")
    for r in after:
        b = before.get((r["system"], r["phase"]))
        if r["system"] in systems or (b and b["repeat"] == "yes"):
            was = b["reason"] or "ok" if b else "(no point)"
            print(f"{r['system']:<9}{r['phase']:<10}{was:<50}{r['reason'] or 'ok'}")

    failed = {s for (s, _), b in before.items() if b["repeat"] == "yes"}
    cmp_dir = OUT / "repeat_compare"
    cmp_dir.mkdir(exist_ok=True)
    for block in sorted({block_of(t) for t in tubes}):
        plot_block(block, [p for p in old if p["block"] == block],
                   [p for p in new if p["block"] == block and p["system"] in systems],
                   failed, two_pe, cmp_dir / f"{block}.png")  # fmt: skip
    print(f"\nWrote {OUT / 'repeat_tielines.csv'}, repeat_list_after.csv, {cmp_dir}/")


if __name__ == "__main__":
    main()
