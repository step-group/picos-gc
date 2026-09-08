# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl"]
# ///
"""Trim the lab prep workbook (aromas_equilibrios_vfinal.xlsx) to the repeat campaign.

Tubes = the tie-lines with a repeat=yes point in out/repeat_list.csv (ternary systems,
binaries as BIN<n> via fill_ternarios.BIN_TO_BLOCK) plus any extra codes on argv. Rows of
every other tube are HIDDEN, not deleted: Lab_DES' volume estimates are cross-sheet
formulas into '2-phenylethanol_DES' and openpyxl would not rewrite them. A re-prepared
tube is a new tie-line, so the printable vial sheets (PLANTILLA_IMPRESION.xlsx's
"Sampling", copied in) keep all four vials (T1 T2 B1 B2) of every kept tube; a block the
template lacks (I - CamEug) is cloned from the last ternary block. Vial codes stay in
the template's hyphen form (C4-T1), which fill_ternarios' CODE_RE accepts.

Run: uv run make_repeat_workbook.py [EXTRA_TUBE ...]   e.g.  uv run make_repeat_workbook.py D2
  -> out/aromas_equilibrios_repeat.xlsx
"""

from __future__ import annotations

import csv
import re
import sys
from copy import copy
from itertools import groupby
from pathlib import Path

import openpyxl

from fill_ternarios import BIN_TO_BLOCK

_ROOT = Path(__file__).resolve().parent
SRC = _ROOT / "aromas_equilibrios_vfinal.xlsx"
LIST = _ROOT / "out" / "repeat_list.csv"
OUT = _ROOT / "out" / "aromas_equilibrios_repeat.xlsx"
_FEED_REF = re.compile(r"'2-phenylethanol_DES'!I(\d+)")
LAB_ROWS, LAB_BIN_HEADER = range(10, 67), 58
FEED_ROWS = range(6, 64)
SAMPLING_SRC = _ROOT / "PLANTILLA_IMPRESION.xlsx"  # sheet "Sampling": printable vial sheets
_BLOCK_TITLE = re.compile(r"^([A-Z]) — ")  # "C — ThyCarvac  (Thymol + Carvacrol)"
# DES prep table (10 g, 1:1 molar): rows 17-25 of datos_des, DES name in B. Sheet1 is
# the 5 g variant: hidden in the output, every batch is made at 10 g to have spare.
DES_TABLES = (("datos_des", range(17, 26), "B"),)
HIDE_SHEETS = ("Sheet1", "Sheet2")  # Sheet2 (binary GC vials) is superseded by Sampling
BIN_V_DES_ML = 4  # Lab_DES F59:F66 (binary rows carry V_DES in F, no density)


def tubes_from_list(path: Path) -> set[str]:
    bin_of = {blk: n for n, blk in BIN_TO_BLOCK.items()}
    tubes = set()
    with path.open() as fh:
        for r in csv.DictReader(fh):
            if r["repeat"] == "yes":
                tubes.add(r["system"] if r["kind"] == "ternary" else f"BIN{bin_of[r['block']]}")
    return tubes


def _backfill(ws, row: int) -> None:
    """Block labels (DES/HBA/HBD in B-D) sit only on a block's first row: copy them down."""
    if ws[f"B{row}"].value is not None:
        return
    src = row - 1
    while ws[f"B{src}"].value is None:
        src -= 1
    for c in "BCD":
        ws[f"{c}{row}"].value = ws[f"{c}{src}"].value


def _hide(ws, rows, keep: set[int]) -> None:
    for r in rows:
        ws.row_dimensions[r].hidden = r not in keep


def _copy_row(ws, src: int, dst: int, ncols: int = 12) -> None:
    for c in range(1, ncols + 1):
        ws.cell(dst, c).value = ws.cell(src, c).value
        ws.cell(dst, c)._style = copy(ws.cell(src, c)._style)
    ws.row_dimensions[dst].height = ws.row_dimensions[src].height


def _copy_sheet(src, wb, title: str):
    """Cross-workbook copy: values + styles attribute-wise (_style ids are per workbook),
    column widths, row heights, landscape fit-to-width. Manual page breaks are dropped:
    they sit on rows that end up hidden."""
    dst = wb.create_sheet(title)
    for row in src.iter_rows():
        for c in row:
            d = dst.cell(c.row, c.column, c.value)
            d.font, d.fill, d.border = copy(c.font), copy(c.fill), copy(c.border)
            d.alignment, d.number_format = copy(c.alignment), c.number_format
    for k, dim in src.column_dimensions.items():
        dst.column_dimensions[k].width = dim.width
    for r, dim in src.row_dimensions.items():
        dst.row_dimensions[r].height = dim.height
    dst.page_setup.orientation = src.page_setup.orientation
    dst.page_setup.fitToWidth, dst.page_setup.fitToHeight = 1, 0
    dst.sheet_properties.pageSetUpPr = copy(src.sheet_properties.pageSetUpPr)
    return dst


def _sampling(wb, tubes: set[str], lab, keep_lab: set[int]) -> None:
    """Bring the print template's Sampling sheet in and keep only the kept tubes' four
    vial rows plus their block's three header rows (title, method line, column header).
    A ternary block the template lacks (I — CamEug) is cloned from the last ternary block
    after the binaries, titled from Lab_DES, with the row formulas re-pointed."""
    ws = _copy_sheet(openpyxl.load_workbook(SAMPLING_SRC)["Sampling"], wb, "Sampling")
    last = ws.max_row
    keep, header, found, tpl = set(), set(), set(), None
    for r in range(1, last + 1):
        a = ws[f"A{r}"].value
        if isinstance(a, str) and (_BLOCK_TITLE.match(a) or a.startswith("BINARIOS")):
            header = {r, r + 1, r + 2}
            tpl = r if _BLOCK_TITLE.match(a) else tpl
        elif a in tubes:
            keep |= header | set(range(r, r + 4))
            found.add(a)
    _hide(ws, range(1, last + 1), keep)
    missing = sorted(tubes - found)
    if bins := [t for t in missing if t.startswith("BIN")]:
        raise ValueError(f"Sampling: no rows for {bins}")
    row = last + 2
    for _letter, group in groupby(missing, key=lambda t: t[0]):
        codes = list(group)
        lab_row = next(r for r in keep_lab if lab[f"A{r}"].value == codes[0])
        des, hba, hbd = (lab[f"{c}{lab_row}"].value for c in "BCD")
        for i in range(3):
            _copy_row(ws, tpl + i, row + i)
        ws[f"A{row}"] = f"{codes[0][0]} — {des}  ({hba} + {hbd})"
        row += 3
        for tube in codes:
            for j, tag in enumerate(("T1", "T2", "B1", "B2")):
                src = tpl + 3 + j  # the template block's first tube: Top 1/2, Bot 1/2
                _copy_row(ws, src, row)
                ws[f"A{row}"] = tube if j == 0 else None
                ws[f"D{row}"] = f"{tube}-{tag}"
                for c in "HIJ":
                    ws[f"{c}{row}"] = re.sub(rf"(?<=[A-L]){src}\b", str(row), ws[f"{c}{src}"].value)
                row += 1


def trim(wb, tubes: set[str]) -> None:
    lab = wb["Lab_DES"]
    known = {lab[f"A{r}"].value for r in LAB_ROWS}
    if missing := tubes - known:
        raise ValueError(f"not in Lab_DES: {sorted(missing)}")
    keep_lab = {r for r in LAB_ROWS if lab[f"A{r}"].value in tubes}
    for r in keep_lab:
        _backfill(lab, r)
    has_bin = any(t.startswith("BIN") for t in tubes)
    _hide(lab, LAB_ROWS, keep_lab | ({LAB_BIN_HEADER} if has_bin else set()))

    des = {lab[f"B{r}"].value for r in keep_lab}  # after back-fill every kept row names it
    for sheet, rows, cols in DES_TABLES:
        ws = wb[sheet]
        if missing := des - {ws[f"{c}{r}"].value for r in rows for c in cols}:
            raise ValueError(f"{sheet}: no prep row for {sorted(missing)}")
        _hide(ws, rows, {r for r in rows if any(ws[f"{c}{r}"].value in des for c in cols)})
    for sheet in HIDE_SHEETS:
        wb[sheet].sheet_state = "hidden"

    feed = wb["2-phenylethanol_DES"]
    keep_feed = set()
    for r in keep_lab:
        if m := _FEED_REF.search(str(lab[f"G{r}"].value or "")):
            keep_feed.add(int(m.group(1)))
    for r in keep_feed:
        _backfill(feed, r)
    _hide(feed, FEED_ROWS, keep_feed)

    _sampling(wb, tubes, lab, keep_lab)


def des_need(tubes: set[str]) -> dict[str, float]:
    """Grams of each DES the kept tubes consume, from Lab_DES' cached estimates
    (ternary rows: V_DES est [I] x rho_DES est [F]; binary rows: 4 mL x that DES' rho)."""
    lab = openpyxl.load_workbook(SRC, data_only=True)["Lab_DES"]
    rows = {r: lab[f"A{r}"].value for r in LAB_ROWS}
    des_of, rho = {}, {}
    for r, code in rows.items():  # block labels only on first rows: carry the DES down
        if lab[f"B{r}"].value:
            block_des = lab[f"B{r}"].value
        if code:
            des_of[r] = block_des
            if r < LAB_BIN_HEADER:
                rho[block_des] = lab[f"F{r}"].value
    need: dict[str, float] = {}
    for r, code in rows.items():
        if code in tubes:
            ml = lab[f"I{r}"].value if r < LAB_BIN_HEADER else BIN_V_DES_ML
            need[des_of[r]] = need.get(des_of[r], 0) + ml * rho.get(des_of[r], 1.0)
    return need


def main() -> None:
    tubes = tubes_from_list(LIST) | set(sys.argv[1:])
    wb = openpyxl.load_workbook(SRC)
    trim(wb, tubes)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    order = sorted(tubes, key=lambda t: (t.startswith("BIN"), t))
    n_tern = sum(not t.startswith("BIN") for t in tubes)
    print(f"{len(tubes)} tubes: {' '.join(order)}; {4 * n_tern} ternary GC vials. Wrote {OUT}")
    print(
        "DES needed (est.): "
        + ", ".join(f"{d} {g:.1f} g" for d, g in sorted(des_need(tubes).items()))
    )


if __name__ == "__main__":
    main()
