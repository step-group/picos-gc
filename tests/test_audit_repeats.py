import pytest

pytest.importorskip("openpyxl")
import openpyxl

from audit_repeats import audit


def _sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bloque Z"
    ws["F2"], ws["G2"], ws["H2"] = 100.0, 100.0, 100.0  # area per %m/m
    ws["M3"], ws["N3"] = "Thymol", "Carvone"
    # Z1: rows 5-6 Superior (both vials, closes), rows 7-8 Inferior (vial 8 never injected)
    # Z2: rows 9-12, no areas at all.
    ws["A5"], ws["A9"] = "Z1", "Z2"
    phases = ["Superior", "Superior", "Inferior", "Inferior"] * 2
    for r, phase in zip(range(5, 13), phases, strict=True):
        ws[f"B{r}"] = phase
        ws[f"D{r}"], ws[f"F{r}"], ws[f"H{r}"] = 10.0, 11.0, 12.0  # dilution x2
        ws[f"U{r}"] = 20.0  # KF 20 % water
    for r in (5, 6, 7):  # 2PE 20 %, HBA 20 %, HBD 20 % (diluted 10 % each) + 20 % water = 0.8
        ws[f"L{r}"], ws[f"M{r}"], ws[f"N{r}"] = 1000.0, 1000.0, 1000.0
    return wb


def test_verdicts():
    by = {(r["system"], r["phase"]): r for r in audit(_sheet())}
    assert by[("Z1", "Superior")]["repeat"] == "no"
    assert by[("Z1", "Superior")]["closure"] == "0.80000"
    assert by[("Z1", "Inferior")]["reason"] == "single_vial"
    assert by[("Z2", "Superior")]["reason"] == "missing_vials"
    assert by[("Z2", "Inferior")]["n_kf"] == 2
    assert len(by) == 4  # no binary block on a synthetic sheet
