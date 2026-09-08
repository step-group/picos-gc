# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl"]
# ///
"""Which LLE tie-line points must be repeated? One row per (system, phase) point.

Reads the *filled* workbook only (no out/ dependency): blank L/M/N areas mean no
chromatogram reached the sheet, i.e. the vial was never injected. Reuses
fill_ternarios' own composition chain (results_rows) so closure/flags here equal
the pipeline's. Binary edge rows 25-28 get the same treatment (fill_ternarios'
binary export carries no closure at all).

Repeat reasons: missing_vials, low_closure, replicate_mismatch,
aqueous_organics_suspect, aqueous_replicate_mismatch. Informational only (in
`flags`, never a repeat): single_vial (one clean vial is accepted), dropped_replicate
(the pipeline already cherry-picked the clean vial, see fill_ternarios.aqueous_keep),
a single KF titration (n_kf).

aqueous_organics_suspect = organic-phase droplets in EVERY kept aqueous vial (terpenes
HBA+HBD above the pair's summed pure-water solubility at 30 °C, see
fill_ternarios.aq_terpene_max; nothing left to cherry-pick).
aqueous_replicate_mismatch = the kept aqueous vials' total organics (2PE+HBA+HBD)
differ > MISMATCH_MAX x without the terpene signature of droplets (a 2PE-only
disagreement). The pipeline's replicate_mismatch is gated on a component > 10 %,
which an aqueous phase never has, so it is blind there. ponytail: no cherry-pick
here — the only reference would be the same tie-line's organic phase via the 2PE
distribution coefficient (40-50 in every clean system), and the one current case
(D1) has a broken organic phase; add that pick when a case with a sound one appears.

Run: uv run audit_repeats.py   ->  out/repeat_list.csv + console table
"""

from __future__ import annotations

import csv
from pathlib import Path

import openpyxl

from fill_ternarios import (
    BIN_PAIRS,
    BIN_TO_BLOCK,
    MISMATCH_MAX,
    ORGANIC_WATER_MAX,
    _num,
    _terp,
    _vial_rows,
    aq_terpene_max,
    aqueous_keep,
    binary_endpoint,
    binary_vials,
    results_rows,
    vial_fractions,
)

_ROOT = Path(__file__).resolve().parent
WB = _ROOT / "Sistemas ternarios_MF_filled.xlsx"
OUT = _ROOT / "out" / "repeat_list.csv"
REPEAT_FLAGS = ("low_closure", "replicate_mismatch")
COLS = [
    "kind", "block", "system", "phase", "hba", "hbd", "codes", "dropped", "n_expected",
    "n_with_areas", "n_vials_used", "n_kf", "closure", "water_src", "aq_terpene_max",
    "aq_ceiling", "aq_organics_ratio", "flags", "repeat", "reason",
]  # fmt: skip
_BIN_CODE = {25: "T1", 26: "T2", 27: "B1", 28: "B2"}  # fill_ternarios._BIN_ROWS inverted


def _codes(prefix: str, reps: list[str], have: list[bool]) -> str:
    """Sample codes to prepare: only the missing ones when some are missing, else all."""
    missing = [c for c, ok in zip(reps, have, strict=True) if not ok]
    return ", ".join(f"{prefix}-{c}" for c in (missing or reps))


def _has_areas(ws, r: int) -> bool:
    return any(_num(ws, f"{c}{r}") is not None for c in "LMN")


def _n_kf(ws, rows) -> int:
    return sum(_num(ws, f"{c}{r}") is not None for r in rows for c in "UV")


def _aqueous_organics(
    kept: list[dict], all_kf: list[float], terp_max: float
) -> tuple[float | None, float | None, list[str]]:
    """(max terpene fraction, max/min total-organics ratio, repeat reasons) over the KEPT
    vials of an aqueous phase. (None, None, []) when the phase is organic or empty."""
    if not all_kf or sum(all_kf) / len(all_kf) < ORGANIC_WATER_MAX or not kept:
        return None, None, []
    terp = [_terp(v) for v in kept]
    tot = [(v["s"] or 0) + t for v, t in zip(kept, terp, strict=True)]
    ratio = (max(tot) / min(tot)) if len(tot) > 1 and min(tot) > 0 else None
    reasons = []
    if min(terp) > terp_max:  # every kept vial carries droplets: nothing to pick
        reasons.append("aqueous_organics_suspect")
    elif ratio is not None and ratio > MISMATCH_MAX:  # 2PE-only disagreement
        reasons.append("aqueous_replicate_mismatch")
    return max(terp), ratio, reasons


def _verdict(row: dict, flags: str) -> dict:
    reason = []
    if row["n_with_areas"] == 0:
        reason.append("missing_vials")
    elif row["n_with_areas"] < row["n_expected"]:
        flags = ";".join(f for f in (flags, "single_vial") if f)
    reason += [f for f in flags.split(";") if f in REPEAT_FLAGS]
    reason += row.pop("aq_reasons", [])
    row["flags"] = flags
    row["repeat"] = "yes" if reason else "no"
    row["reason"] = ";".join(reason)
    return row


def _ternary(ws, block: str) -> list[dict]:
    groups: dict[tuple[int, str], list[int]] = {}
    for row, sysnum, ph, _rep, phase in _vial_rows(ws):
        if phase:
            groups.setdefault((sysnum, ph), []).append(row)
    recs = []
    for (sysnum, ph), rows in groups.items():
        for r in rows:
            if not _has_areas(ws, r):
                continue
            d, f, h = _num(ws, f"D{r}"), _num(ws, f"F{r}"), _num(ws, f"H{r}")
            rec = {
                "row": r, "sysnum": sysnum, "ph": ph, "phase": ws[f"B{r}"].value,
                "L": _num(ws, f"L{r}"), "M": _num(ws, f"M{r}"), "N": _num(ws, f"N{r}"),
                "I": (f - d) if None not in (f, d) else None,
                "K": (h - d) if None not in (h, d) else None,
                "U": _num(ws, f"U{r}"), "V": _num(ws, f"V{r}"),
            }  # fmt: skip
            recs.append(rec)
    # results_rows: [block, system, phase, ..., closure(9), n_vials(10), flags(11), src(12)]
    computed = {(r[1], r[2]): r for r in results_rows(ws, block, recs)}
    f2, g2, h2 = _num(ws, "F2"), _num(ws, "G2"), _num(ws, "H2")
    terp_max = aq_terpene_max(ws["M3"].value, ws["N3"].value)
    out = []
    for (sysnum, ph), rows in sorted(groups.items()):
        system, phase = f"{block}{sysnum}", "Superior" if ph == "T" else "Inferior"
        c = computed.get((system, phase))
        vials = [
            v
            for r in recs
            if (r["sysnum"], r["ph"]) == (sysnum, ph)
            and (v := vial_fractions(r, f2, g2, h2)) is not None
        ]
        kept = aqueous_keep(vials, terp_max)
        all_kf = [c for v in vials for c in v["kf"]]
        terp, ratio, aq_reasons = _aqueous_organics(kept, all_kf, terp_max)
        have = [_has_areas(ws, r) for r in rows]
        reps = [f"{ph}{i}" for i in range(1, len(rows) + 1)]
        dropped = [v["raw"]["row"] for v in vials if v not in kept]
        row = {
            "kind": "ternary", "block": block, "system": system, "phase": phase,
            "hba": ws["M3"].value, "hbd": ws["N3"].value,
            "codes": _codes(system, reps, have),
            "dropped": ", ".join(f"{system}-{reps[rows.index(r)]}" for r in dropped),
            "n_expected": len(rows),
            "n_with_areas": sum(have),
            "n_vials_used": c[10] if c else 0,
            "n_kf": _n_kf(ws, rows),
            "closure": c[9] if c else "",
            "water_src": c[12] if c else "",
            "aq_terpene_max": f"{terp:.5f}" if terp is not None else "",
            "aq_ceiling": f"{terp_max:.5f}" if terp is not None else "",
            "aq_organics_ratio": f"{ratio:.2f}" if ratio is not None else "",
            "aq_reasons": aq_reasons,
        }  # fmt: skip
        out.append(_verdict(row, c[11] if c else ""))
    return out


def _binary(ws, block: str) -> list[dict]:
    if ws["A25"].value != "Bin":
        return []
    g2, h2 = _num(ws, "G2"), _num(ws, "H2")
    hba, hbd = ws["M3"].value, ws["N3"].value
    terp_max = aq_terpene_max(hba, hbd)
    binnum = next((b for b, blk in BIN_TO_BLOCK.items() if blk == block), None)
    prefix = f"BIN{binnum}" if binnum else f"{block}-bin"
    out = []
    for _row, reps in BIN_PAIRS:
        vials = binary_vials(ws, reps, g2, h2)
        point, closure, src, flags, kept = binary_endpoint(vials, terp_max)
        all_kf = [c for v in vials for c in v["kf"]]
        terp, ratio, aq_reasons = _aqueous_organics(kept, all_kf, terp_max)
        have = [_num(ws, f"M{x}") is not None or _num(ws, f"N{x}") is not None for x in reps]
        dropped = [v["raw"]["row"] for v in vials if v not in kept]
        r = {
            "kind": "binary", "block": block, "system": f"{block}-bin",
            "phase": "" if point is None else ("aqueous" if point[2] > 0.5 else "organic"),
            "hba": hba, "hbd": hbd,
            "codes": _codes(prefix, [_BIN_CODE[x] for x in reps], have),
            "dropped": ", ".join(f"{prefix}-{_BIN_CODE[x]}" for x in dropped),
            "n_expected": len(reps),
            "n_with_areas": sum(have),
            "n_vials_used": len(kept),
            "n_kf": _n_kf(ws, reps),
            "closure": f"{closure:.5f}" if closure is not None else "",
            "water_src": src,
            "aq_terpene_max": f"{terp:.5f}" if terp is not None else "",
            "aq_ceiling": f"{terp_max:.5f}" if terp is not None else "",
            "aq_organics_ratio": f"{ratio:.2f}" if ratio is not None else "",
            "aq_reasons": aq_reasons,
        }  # fmt: skip
        out.append(_verdict(r, ";".join(flags)))
    return out


def audit(wb) -> list[dict]:
    """All (system, phase) points of a data_only workbook, with a repeat verdict each."""
    rows = []
    for sheet in wb.sheetnames:
        ws, block = wb[sheet], sheet.split()[-1]
        rows += _ternary(ws, block) + _binary(ws, block)
    return rows


def main() -> None:
    rows = audit(openpyxl.load_workbook(WB, data_only=True))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=COLS)
        w.writeheader()
        w.writerows(rows)
    rep = [r for r in rows if r["repeat"] == "yes"]
    print(
        f"{'point':<10}{'phase':<10}{'vials':<7}{'closure':<10}{'aq_terp':<9}{'aq_ratio':<10}reason"
    )
    for r in rep:
        vials = f"{r['n_with_areas']}/{r['n_expected']}"
        print(
            f"{r['system']:<10}{r['phase']:<10}{vials:<7}{r['closure']:<10}"
            f"{r['aq_terpene_max']:<9}{r['aq_organics_ratio']:<10}{r['reason']}"
        )
    print(f"\n{len(rep)} of {len(rows)} points need repeating. Wrote {OUT}")


if __name__ == "__main__":
    main()
