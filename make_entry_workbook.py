# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl"]
# ///
"""Data-entry workbook for the repeat campaign: gravimetry of every vial + preliminary KF.

One row per vial (T1 T2 B1 B2) of every repeat tube (out/repeat_list.csv, as in
make_repeat_workbook) plus the 2PE-water binary. Masses are typed for every vial; KF only
on the ORGANIC phase (aqueous water is by difference) and only for the tubes in FRESH_KF --
the others show their campaign-1 KF from Sistemas ternarios_MF.xlsx (U/V), locked. The
sheet is protected without a password, so only the yellow cells take input. The row under
the header names the master workbook's column for each value, to copy it across.

Run: uv run make_entry_workbook.py [EXTRA_TUBE ...]   -> out/repeat_entry.xlsx
Re-running keeps whatever was typed into an existing out/repeat_entry.xlsx (matched by vial
code, previous file kept as .bak.xlsx) and refuses to save if a typed value has no cell left.
"""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation

from fill_ternarios import _BIN_ROWS, BIN_TO_BLOCK, ORGANIC_WATER_MAX, WB_IN, _vial_rows
from make_repeat_workbook import LIST, SAMPLE_UL, round_args, tubes_from_list

OUT = Path(__file__).resolve().parent / "out" / "repeat_entry.xlsx"
# Organic KF re-measured by the user in the repeat campaign (2026-09-22); every other
# repeat tube keeps its campaign-1 organic KF.
FRESH_KF = {"C4", "C5", "D1", "E1", "I1", "2PE"}
# Organic phase where campaign-1 KF cannot decide it. 2PE-water has no history: the
# 2PE-rich organic phase is Inferior in all eight blocks, so bottom. Flip to "T" if it floats.
ORGANIC_OVERRIDE = {"2PE": "B"}
# Re-prepared although repeat_list.csv has no repeat=yes point for them (D2: dropped replicate).
EXTRA_TUBES = {"2PE", "D2"}
VIALS = (("T", 1), ("T", 2), ("B", 1), ("B", 2))
PHASE = {"T": "Superior", "B": "Inferior"}

HEAD = ["Tube", "Vial code", "Phase", "Role", "Vial, g", "Sample, µL", "Vial + sample, g",
        "IPA, µL", "Vial + sample + IPA, g", "Sample, g", "IPA, g", "Dilution factor",
        "KF 1, % m/m", "KF 2, % m/m", "KF status", "Notes"]  # fmt: skip
MASTER_COL = {"E": "D", "F": "E", "G": "F", "H": "G", "I": "H", "J": "I", "K": "J", "M": "U", "N": "V"}
WIDTH = dict(zip("ABCDEFGHIJKLMNOP", (7, 11, 10, 9, 9, 9, 12, 8, 13, 9, 9, 9, 10, 10, 22, 28), strict=True))
HEAD_ROW, LETTER_ROW, EXAMPLE_ROW, FIRST = 6, 7, 8, 9

ARIAL = Font(name="Arial", size=10)
INPUT = PatternFill("solid", fgColor="FFFF00")
REUSED = PatternFill("solid", fgColor="D9D9D9")
AQUEOUS = PatternFill("solid", fgColor="808080")
UNLOCKED = Protection(locked=False)
TUBE_TOP = Border(top=Side(style="thin"))


def _sheet_of(tube: str) -> str:
    return f"Bloque {BIN_TO_BLOCK[int(tube[3:])] if tube.startswith('BIN') else tube[0]}"


def campaign1_kf(wb, tube: str) -> dict[tuple[str, int], tuple]:
    """(phase, rep) -> (KF1, KF2) of the tube's campaign-1 vials, read from U/V."""
    ws = wb[_sheet_of(tube)]
    if tube.startswith("BIN"):
        rows = dict(_BIN_ROWS)
    else:
        rows = {(ph, rep): row for row, n, ph, rep, _ in _vial_rows(ws) if n == int(tube[1:])}
    return {key: (ws[f"U{row}"].value, ws[f"V{row}"].value) for key, row in rows.items()}


def organic_phase(tube: str, kf: dict) -> str:
    if tube in ORGANIC_OVERRIDE:
        return ORGANIC_OVERRIDE[tube]
    for ph in "TB":
        vals = [v for rep in (1, 2) for v in kf.get((ph, rep), ()) if isinstance(v, int | float)]
        if vals and sum(vals) / len(vals) < 100 * ORGANIC_WATER_MAX:
            return ph
    raise ValueError(f"{tube}: no campaign-1 KF below {100 * ORGANIC_WATER_MAX:.0f} % in either phase")


def _order(tube: str) -> tuple:
    if tube == "2PE":
        return ("Z", 0, 0)
    if tube.startswith("BIN"):
        return (BIN_TO_BLOCK[int(tube[3:])], 1, 0)
    return (tube[0], 0, int(tube[1:]))


def _row(ws, r: int, tube: str, code: str, ph: str, role: str) -> None:
    ws.append([tube, code, PHASE[ph], role, None, SAMPLE_UL, None, f"=1000-F{r}", None,
               f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(G{r})),G{r}-E{r},"")',
               f'=IF(AND(ISNUMBER(G{r}),ISNUMBER(I{r})),I{r}-G{r},"")',
               f'=IF(AND(ISNUMBER(J{r}),ISNUMBER(K{r})),IF(J{r}>0,(J{r}+K{r})/J{r},""),"")'])  # fmt: skip
    for c in "EFGIP":
        ws[f"{c}{r}"].fill, ws[f"{c}{r}"].protection = INPUT, UNLOCKED
        ws[f"{c}{r}"].font = Font(name="Arial", size=10, color="0000FF")
    for c in "EGIJK":
        ws[f"{c}{r}"].number_format = "0.0000"
    ws[f"L{r}"].number_format = "0.00"


def build(tubes: set[str], fresh: set[str] = FRESH_KF) -> openpyxl.Workbook:
    master = openpyxl.load_workbook(WB_IN)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Entry"
    ws.append(["Repeat campaign: vial gravimetry and preliminary Karl-Fischer"])
    ws.append([None, "Type here (the only editable cells)"])
    ws.append([None, "Campaign-1 KF reused, locked"])
    ws.append([None, "Aqueous phase: water by difference, no KF"])
    ws.append([None, "Sheet protected without password (Review > Unprotect sheet). Row 7 names the"
               " column of Sistemas ternarios_MF.xlsx each value goes to."])  # fmt: skip
    ws.append(HEAD)
    ws.append([MASTER_COL.get(c) for c in "ABCDEFGHIJKLMNOP"])
    for cell, fill in (("A2", INPUT), ("A3", REUSED), ("A4", AQUEOUS)):
        ws[cell].fill = fill

    _row(ws, EXAMPLE_ROW, "example", "A1-T1", "T", "organic")
    for c, v in (("E", 2.7385), ("G", 2.9211), ("I", 3.4876), ("M", 1.93), ("O", "fresh (example)")):
        ws[f"{c}{EXAMPLE_ROW}"] = v
    for cell in ws[EXAMPLE_ROW]:
        cell.fill, cell.protection = PatternFill(), Protection(locked=True)
        cell.font = Font(name="Arial", size=10, italic=True, color="808080")

    kf_input, mass_input = [], []
    r = FIRST
    for tube in sorted(tubes, key=_order):
        kf = {} if tube == "2PE" else campaign1_kf(master, tube)
        org = organic_phase(tube, kf)
        for i, (ph, rep) in enumerate(VIALS):
            role = "organic" if ph == org else "aqueous"
            _row(ws, r, tube, f"{tube}-{ph}{rep}", ph, role)
            old = kf.get((ph, rep), (None, None))
            if role == "aqueous":
                ws[f"O{r}"] = "n/a: by difference"
                for c in "MN":
                    ws[f"{c}{r}"].fill = AQUEOUS
            elif tube in fresh:
                ws[f"O{r}"] = "fresh: type the new KF"
                for c, v in zip("MN", old, strict=True):
                    ws[f"{c}{r}"].fill, ws[f"{c}{r}"].protection = INPUT, UNLOCKED
                    ws[f"{c}{r}"].font = Font(name="Arial", size=10, color="0000FF")
                    if v is not None:
                        ws[f"{c}{r}"].comment = Comment(f"Campaign 1: {v} %", "make_entry_workbook")
                kf_input.append(f"M{r}:N{r}")
            else:
                ws[f"O{r}"] = "reused: campaign 1"
                for c, v in zip("MN", old, strict=True):
                    ws[f"{c}{r}"], ws[f"{c}{r}"].fill = v, REUSED
            if i == 0:
                for cell in ws[r]:
                    cell.border = TUBE_TOP
            mass_input.append(f"E{r}:G{r} I{r}")
            r += 1

    for row in ws.iter_rows():
        for cell in row:
            if cell.font.name != "Arial":  # not styled above
                cell.font = ARIAL
    for cell in ws[1]:
        cell.font = Font(name="Arial", size=12, bold=True)
    for cell in ws[HEAD_ROW]:
        cell.font = Font(name="Arial", size=10, bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for cell in ws[LETTER_ROW]:
        cell.alignment = Alignment(horizontal="center")
    for c, w in WIDTH.items():
        ws.column_dimensions[c].width = w
    ws.row_dimensions[HEAD_ROW].height = 40
    for rng, lo, hi in ((" ".join(mass_input), 0, 100), (" ".join(kf_input), 0, 100)):
        dv = DataValidation("decimal", formula1=str(lo), formula2=str(hi), sqref=rng)
        dv.error, dv.showErrorMessage = f"A number between {lo} and {hi}", True
        ws.add_data_validation(dv)
    ws.freeze_panes = f"C{EXAMPLE_ROW}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HEAD_ROW}:{LETTER_ROW}"
    ws.protection.sheet = True
    return wb


def typed(ws) -> dict[str, dict[str, object]]:
    """Vial code -> {column: value} of every unlocked (= input) cell holding more than its default."""
    out = {}
    for r in range(FIRST, ws.max_row + 1):
        vals = {c.column_letter: c.value for c in ws[r] if not c.protection.locked and c.value is not None}
        if vals.get("F") == SAMPLE_UL:
            del vals["F"]
        if vals:
            out[ws[f"B{r}"].value] = vals
    return out


def carry(ws, inputs: dict[str, dict[str, object]]) -> list[str]:
    """Write `inputs` back into a fresh sheet; return what has no input cell left to go to."""
    row = {ws[f"B{r}"].value: r for r in range(FIRST, ws.max_row + 1)}
    lost = []
    for code, vals in inputs.items():
        for c, v in vals.items():
            if code not in row or ws[f"{c}{row[code]}"].protection.locked:
                lost.append(f"{code} {c}={v}")
            else:
                ws[f"{c}{row[code]}"] = v
    return lost


def main() -> None:
    tubes, out = round_args(sys.argv[1:], OUT)
    if tubes is None:  # round 1
        tubes, fresh = tubes_from_list(LIST) | EXTRA_TUBES | set(sys.argv[1:]), FRESH_KF
    else:  # a later round's tubes are new tie-lines: every organic phase gets its own KF
        fresh = tubes
    out.parent.mkdir(exist_ok=True)
    wb = build(tubes, fresh)
    if out.exists():  # the user types into out: a rebuild keeps every typed cell
        lost = carry(wb.active, typed(openpyxl.load_workbook(out).active))
        if lost:
            sys.exit(f"refusing to overwrite {out.name}, these inputs have no cell left:\n  " + "\n  ".join(lost))
        shutil.copy2(out, out.with_suffix(".bak.xlsx"))
    wb.save(out)
    print(f"{out}: {len(tubes)} tubes, {4 * len(tubes)} vials ({', '.join(sorted(tubes, key=_order))})")


if __name__ == "__main__":
    main()
